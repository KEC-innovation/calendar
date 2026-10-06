#!/usr/bin/env python3
"""Audit and idempotently import the private KEC legacy workbook set.

Dry-run is the default and writes only aggregate, non-PII reports. Applying needs
DATABASE_URL and an explicit decision about rows detected as tests. Source files,
calendar identifiers, personal data, and answer keys are never copied into the repo.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import uuid
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

NAMESPACE = uuid.UUID("ed2dba97-0180-4b96-ac41-8623c01e48a7")
NEPAL = ZoneInfo("Asia/Kathmandu")

EQUIPMENT_CERT = {
    "Laser Cutter (Katrina)": "laser-cutting",
    "3D Scanner (Raptor)": "3d-scanning",
    "Anycubic Kobra 3 (Nagini)": "3d-printing",
    "Anycubic Neo (Niro)": "3d-printing",
    "Bambu A1 (1)": "3d-printing",
    "Bambu A1 (2)": "3d-printing",
    "Electronic Station 1": "electronics-workstation",
    "Electronic Station 2": "electronics-workstation",
}
TRAINING_SLUG = {"Laser Cutting": "laser-cutting", "3D Printing": "3d-printing"}
QUIZ_SOURCE_FILES = {
    "Laser Cutting": "KEC Makerspace_Laser Cutting Quiz.xlsx",
    "3D Printing": "KEC Makerspace_3d Printing Quiz.xlsx",
}


def stable_id(kind: str, key: str) -> str:
    return str(uuid.uuid5(NAMESPACE, f"{kind}:{key.strip().casefold()}"))


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def email(value: Any) -> str:
    return clean(value).lower()


def truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return clean(value).lower() in {"true", "yes", "1", "active", "granted"}


def as_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(clean(value)))
    except (TypeError, ValueError):
        return default


def iso_datetime(value: Any, fallback: datetime | None = None) -> str:
    if isinstance(value, datetime):
        current = value if value.tzinfo else value.replace(tzinfo=NEPAL)
        return current.astimezone(timezone.utc).isoformat()
    return (fallback or datetime(2000, 1, 1, tzinfo=timezone.utc)).isoformat()


def category(value: Any) -> str:
    normalized = clean(value).casefold()
    if "kec" in normalized and "student" in normalized:
        return "kec_student"
    if "kec" in normalized and "staff" in normalized:
        return "kec_staff"
    if "college" in normalized or "student" in normalized:
        return "other_college_student"
    if "member" in normalized:
        return "member_non_kec"
    return "business_external"


def is_test_record(name: Any, address: Any, *values: Any) -> bool:
    haystack = " ".join(clean(item).casefold() for item in (name, address, *values))
    address_text = email(address)
    return (
        bool(re.search(r"(^|\W)test(\W|$)", haystack))
        or address_text.endswith("@test.com")
        or address_text.endswith("@123.com")
        or address_text.endswith("@01.com")
        or address_text in {"test@test.com", "test@123.com", "test01@01.com"}
    )


def fingerprint(value: str) -> str:
    return hashlib.sha256(value.strip().casefold().encode()).hexdigest()[:12]


def find_one(source_dir: Path, filename: str) -> Path:
    matches = list(source_dir.rglob(filename))
    if len(matches) != 1:
        raise ValueError(f"Expected one {filename!r}; found {len(matches)}.")
    return matches[0]


def rows_by_header(workbook_path: Path, sheet_name: str) -> list[tuple[int, dict[str, Any]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing sheet {sheet_name!r} in {workbook_path.name}.")
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    return [(number, dict(zip(headers, values))) for number, values in enumerate(iterator, 2)]


def sheet_headers(workbook_path: Path, sheet_name: str) -> set[str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing sheet {sheet_name!r} in {workbook_path.name}.")
    return {clean(value) for value in next(workbook[sheet_name].iter_rows(values_only=True))}


def parse_quiz_bank(code_path: Path) -> dict[str, list[dict[str, Any]]]:
    source = code_path.read_text(encoding="utf-8")
    match = re.search(r"const\s+QUIZ_BANK\s*=\s*(\{.*?\n\});", source, flags=re.DOTALL)
    if not match:
        raise ValueError("QUIZ_BANK was not found in the current Code.gs.")
    bank = json.loads(match.group(1))
    if set(bank) != {"Laser Cutting", "3D Printing"}:
        raise ValueError("Unexpected quiz names in QUIZ_BANK.")
    for quiz_name, questions in bank.items():
        if len(questions) != 20:
            raise ValueError(f"{quiz_name} has {len(questions)} questions; expected 20.")
        for question in questions:
            if question.get("correct") not in question.get("options", []):
                raise ValueError(f"{quiz_name} {question.get('id')} has an invalid correct answer.")
            if len(set(question.get("options", []))) != len(question.get("options", [])):
                raise ValueError(f"{quiz_name} {question.get('id')} has duplicate options.")
    return bank


def quiz_response_rows(path: Path, training: str, bank: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    headers = sheet_headers(path, "Sheet1")
    missing_prompts = [question["text"] for question in bank if question["text"] not in headers]
    if missing_prompts:
        raise ValueError(f"{path.name} is missing {len(missing_prompts)} questions from the current private bank.")
    rows = rows_by_header(path, "Sheet1")
    response_by_id: dict[str, dict[str, Any]] = {}
    for _, row in rows:
        source_id = clean(row.get("Id"))
        if not source_id:
            continue
        if training == "Laser Cutting":
            participant_name, participant_email = row.get("Name1"), row.get("Email1")
        else:
            participant_name = row.get("Name1") or row.get("Name")
            participant_email = row.get("Email Adress") or row.get("Email")
        answers = []
        for question in bank:
            selected = clean(row.get(question["text"]))
            points = as_int(row.get(f"Points - {question['text']}"))
            if selected:
                answers.append({"legacy_question_id": question["id"], "selected_label": selected, "was_correct": points > 0, "matched_option": selected in question["options"]})
        response_by_id[source_id] = {
            "name": clean(participant_name), "email": email(participant_email),
            "started_at": iso_datetime(row.get("Start time")), "completed_at": iso_datetime(row.get("Completion time")),
            "answers": answers,
        }
    return response_by_id


def find_existing_email(people: dict[str, dict[str, Any]], name: Any, roll: Any) -> str:
    wanted_roll = clean(roll).casefold()
    if wanted_roll:
        matches = [address for address, person in people.items() if clean(person["roll_number"]).casefold() == wanted_roll]
        if len(matches) == 1:
            return matches[0]
    # A name alone is not an identity key. A conflicting roll must never fall
    # through to a same-name person and transfer their training/access.
    return ""


def add_person(people: dict[str, dict[str, Any]], issues: list[dict[str, str]], *, address: Any, name: Any, roll: Any, phone: Any, user_type: Any, organization: Any, source_key: str, active: bool = True) -> str | None:
    normalized_email = email(address)
    invalid_address = not bool(re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", normalized_email))
    if invalid_address:
        issues.append({"kind": "invalid_email", "source": source_key, "subject": fingerprint(clean(address) or source_key)})
        normalized_email = f"legacy-missing-{fingerprint(source_key)}@invalid.local"
    normalized_roll = clean(roll)
    current = people.get(normalized_email)
    if current:
        if source_key.startswith("authorized-users:"):
            issues.append({"kind": "duplicate_email", "source": source_key, "subject": fingerprint(normalized_email)})
        if current["roll_number"] and normalized_roll and current["roll_number"].casefold() != normalized_roll.casefold():
            issues.append({"kind": "roll_conflict", "source": source_key, "subject": fingerprint(normalized_email)})
        if not current["phone"] and clean(phone):
            current["phone"] = clean(phone)
        return current["id"]
    person = {
        "id": stable_id("person", normalized_email), "email": normalized_email,
        "full_name": clean(name) or "Legacy participant", "roll_number": normalized_roll or None,
        "phone": clean(phone) or None, "category": category(user_type), "organization": clean(organization) or None,
        "active": active and not invalid_address, "booking_privilege_active": active and not invalid_address,
        "safety_training_status": "unknown", "waiver_status": "unknown", "minor_status": "unknown",
        "migration_review_required": True, "priority_rank": 10 if category(user_type) in {"kec_student", "kec_staff"} else 30,
        "legacy_source_key": source_key,
        "metadata": {"migration": "legacy_workbook", "compliance_not_present_in_source": True, "synthetic_email_for_missing_source_value": invalid_address},
    }
    people[normalized_email] = person
    if person["category"] == "kec_student" and not normalized_roll:
        issues.append({"kind": "missing_roll", "source": source_key, "subject": fingerprint(normalized_email)})
    return person["id"]


def build_model(source_dir: Path, include_tests: bool) -> tuple[dict[str, Any], dict[str, Any], list[dict[str, str]]]:
    main_path = find_one(source_dir, "KEC Makerspace Booking System.xlsx")
    code_path = find_one(source_dir, "Code.gs")
    bank = parse_quiz_bank(code_path)
    issues: list[dict[str, str]] = []
    test_rows: list[dict[str, str]] = []
    people: dict[str, dict[str, Any]] = {}
    permissions: set[tuple[str, str]] = set()

    user_rows = [(number, row) for number, row in rows_by_header(main_path, "Authorized Users") if email(row.get("Email"))]
    for number, row in user_rows:
        test = is_test_record(row.get("Full Name"), row.get("Email"))
        if test:
            test_rows.append({"kind": "authorized_user", "source": str(number), "subject": fingerprint(email(row.get("Email")))})
            if not include_tests:
                continue
        person_id = add_person(people, issues, address=row.get("Email"), name=row.get("Full Name"), roll=row.get("Roll No"), phone=row.get("Phone"), user_type=row.get("User Type"), organization=row.get("Organization / Company"), source_key=f"authorized-users:{number}", active=truthy(row.get("Active")))
        if not person_id:
            continue
        for legacy_name, cert_slug in EQUIPMENT_CERT.items():
            if truthy(row.get(legacy_name)):
                permissions.add((email(row.get("Email")), cert_slug))

    response_lookup: dict[tuple[str, str], dict[str, Any]] = {}
    quiz_validation: dict[str, Any] = {}
    for training, filename in QUIZ_SOURCE_FILES.items():
        source_rows = quiz_response_rows(find_one(source_dir, filename), training, bank[training])
        response_lookup.update({(training, source_id): value for source_id, value in source_rows.items()})
        prompts = {question["text"] for question in bank[training]}
        quiz_validation[training] = {"questions": len(bank[training]), "responses": len(source_rows), "bank_questions_mapped_to_export": len(prompts)}

    attempts: list[dict[str, Any]] = []
    certifications_by_person_type: dict[tuple[str, str], dict[str, Any]] = {}
    training_rows = [(number, row) for number, row in rows_by_header(main_path, "Training Results") if clean(row.get("Training"))]
    for number, row in training_rows:
        training = clean(row.get("Training"))
        source_id = clean(row.get("Source Response ID"))
        response = response_lookup.get((training, source_id), {})
        participant_email = email(row.get("Email"))
        if "@" not in participant_email:
            participant_email = email(response.get("email"))
        if "@" not in participant_email:
            participant_email = find_existing_email(people, row.get("Name") or response.get("name"), row.get("Roll No / External ID"))
        test = is_test_record(row.get("Name"), participant_email, row.get("Notes"))
        if test:
            test_rows.append({"kind": "training_result", "source": str(number), "subject": fingerprint(participant_email or str(number))})
            if not include_tests:
                continue
        person_id = add_person(people, issues, address=participant_email, name=row.get("Name") or response.get("name"), roll=row.get("Roll No / External ID"), phone=row.get("Phone"), user_type=row.get("User Type"), organization=row.get("Organization / Company"), source_key=f"training-person:{training}:{source_id}")
        quiz_slug = TRAINING_SLUG.get(training)
        if not person_id or not quiz_slug:
            issues.append({"kind": "unmapped_training", "source": str(number), "subject": fingerprint(training)})
            continue
        resolved_email = next((address for address, person in people.items() if person["id"] == person_id), participant_email)
        started_at = clean(response.get("started_at")) or iso_datetime(row.get("Granted At"))
        completed_at = clean(response.get("completed_at")) or iso_datetime(row.get("Granted At"), datetime.fromisoformat(started_at) + timedelta(minutes=1))
        score, max_score, pass_mark = as_int(row.get("Score")), as_int(row.get("Max Score"), 20), as_int(row.get("Pass Mark"), 16)
        passed = clean(row.get("Result")).upper() == "PASS" and score >= pass_mark
        attempt_key = f"training-result:{training}:{source_id or number}"
        attempt = {
            "id": stable_id("attempt", attempt_key), "attempt_reference": f"LEGACY-{quiz_slug.upper()}-{source_id or number}",
            "quiz_slug": quiz_slug, "participant_id": person_id, "trainer_name_snapshot": "Legacy form import",
            "status": "submitted", "started_at": started_at, "expires_at": iso_datetime(None, datetime.fromisoformat(started_at) + timedelta(minutes=10)),
            "submitted_at": completed_at, "score": score, "max_score": max_score, "pass_mark": pass_mark, "passed": passed,
            "legacy_source_key": attempt_key, "answers": response.get("answers", []),
            "source_metadata": {"source_response_id": source_id, "legacy_result": clean(row.get("Result")), "legacy_grant_status": clean(row.get("Grant Status"))},
        }
        attempts.append(attempt)
        unmatched = sum(1 for answer in attempt["answers"] if not answer["matched_option"])
        if unmatched:
            issues.append({"kind": "answer_not_in_current_bank", "source": attempt_key, "subject": str(unmatched)})
        granted = clean(row.get("Grant Status")).upper() == "GRANTED" and passed
        if granted:
            cert_key = (resolved_email, quiz_slug)
            certifications_by_person_type[cert_key] = {
                "id": stable_id("certification", f"{resolved_email}:{quiz_slug}"), "person_id": person_id,
                "certification_slug": quiz_slug, "status": "active", "issued_at": iso_datetime(row.get("Granted At"), datetime.fromisoformat(completed_at)),
                "source_attempt_id": attempt["id"], "source_kind": "legacy_training_result",
                "reason": "Migrated from a granted legacy training result.", "legacy_source_key": f"training-cert:{training}:{source_id or number}",
                "metadata": {"legacy_equipment": clean(row.get("Eligible Equipment"))},
            }
        elif passed:
            issues.append({"kind": "passed_without_grant", "source": attempt_key, "subject": fingerprint(resolved_email)})

    for participant_email, cert_slug in sorted(permissions):
        person = people.get(participant_email)
        if not person or (participant_email, cert_slug) in certifications_by_person_type:
            continue
        certifications_by_person_type[(participant_email, cert_slug)] = {
            "id": stable_id("certification", f"{participant_email}:{cert_slug}"), "person_id": person["id"],
            "certification_slug": cert_slug, "status": "active", "issued_at": None, "source_attempt_id": None,
            "source_kind": "legacy_permission", "reason": "Migrated from an explicit active equipment permission; issue date was not present.",
            "legacy_source_key": f"permission-cert:{participant_email}:{cert_slug}", "metadata": {"evidence": "explicit_equipment_permission_in_supplied_people_list"},
        }

    # Recognize existing equipment training without inventing a separate general
    # safety induction, signed waiver, or age. Missing contact details are a
    # records-review task, never a reason to ask someone to repeat a passed quiz.
    for person in people.values():
        recognized = sorted(cert["certification_slug"] for cert in certifications_by_person_type.values() if cert["person_id"] == person["id"])
        passed_types = sorted({attempt["quiz_slug"] for attempt in attempts if attempt["participant_id"] == person["id"] and attempt["passed"]})
        person["metadata"]["recognized_equipment_training"] = recognized
        person["metadata"]["passed_training"] = passed_types
        person["metadata"]["training_retake_required"] = False if recognized or passed_types else None

    legacy_equipment: list[dict[str, Any]] = []
    equipment_rows = [(number, row) for number, row in rows_by_header(main_path, "Equipment") if clean(row.get("Equipment"))]
    for number, row in equipment_rows:
        name = clean(row.get("Equipment"))
        legacy_equipment.append({"legacy_name": name, "calendar_id": clean(row.get("Calendar ID")) or None, "active": truthy(row.get("Active")), "external_allowed": truthy(row.get("External Clients Allowed")), "source_row": number})

    bookings: list[dict[str, Any]] = []
    booking_rows = [(number, row) for number, row in rows_by_header(main_path, "Bookings") if clean(row.get("Booking ID"))]
    future_slots: dict[str, list[tuple[datetime, datetime]]] = {}
    migration_time = datetime.now(timezone.utc)
    for number, row in booking_rows:
        participant_email = email(row.get("Email"))
        test = is_test_record(row.get("Full Name"), participant_email, row.get("Reason"), row.get("Booking ID"))
        if test:
            test_rows.append({"kind": "booking", "source": str(number), "subject": fingerprint(participant_email or str(number))})
            if not include_tests:
                continue
        person_id = add_person(people, issues, address=participant_email, name=row.get("Full Name"), roll=row.get("Roll No"), phone=row.get("Phone"), user_type=row.get("User Type"), organization=row.get("Organization / Company"), source_key=f"booking-person:{clean(row.get('Booking ID'))}")
        if not person_id:
            continue
        starts_at = datetime.fromisoformat(iso_datetime(row.get("Start")))
        ends_at = datetime.fromisoformat(iso_datetime(row.get("End")))
        original_status = clean(row.get("Status")).lower() or "confirmed"
        status = original_status if original_status in {"confirmed", "checked_in", "completed", "cancelled", "no_show"} else "confirmed"
        normalized_reason = None
        if status in {"confirmed", "checked_in"} and ends_at <= migration_time:
            status, normalized_reason = "completed", "past_active_to_completed"
        slots = future_slots.setdefault(clean(row.get("Equipment")), [])
        if status in {"confirmed", "checked_in"} and any(starts_at < existing_end and ends_at > existing_start for existing_start, existing_end in slots):
            status, normalized_reason = "cancelled", "future_overlap_to_cancelled"
            issues.append({"kind": "booking_overlap", "source": clean(row.get("Booking ID")), "subject": fingerprint(clean(row.get("Equipment")))})
        if status in {"confirmed", "checked_in"}:
            slots.append((starts_at, ends_at))
        bookings.append({
            "id": stable_id("booking", clean(row.get("Booking ID"))), "booking_reference": clean(row.get("Booking ID")),
            "person_id": person_id, "legacy_equipment": clean(row.get("Equipment")), "status": status,
            "starts_at": starts_at.isoformat(), "ends_at": ends_at.isoformat(), "contact_name": clean(row.get("Full Name")),
            "contact_email": participant_email, "contact_phone": clean(row.get("Phone")) or "Not recorded",
            "contact_roll_number": clean(row.get("Roll No")) or None, "contact_organization": clean(row.get("Organization / Company")) or None,
            "purpose": clean(row.get("Reason")) or None, "cancelled_at": migration_time.isoformat() if status == "cancelled" else None,
            "completed_at": ends_at.isoformat() if status == "completed" else None,
            "calendar_event_id": clean(row.get("Calendar Event ID")) or None, "calendar_sync_status": "synced" if clean(row.get("Calendar Event ID")) else "not_configured",
            "legacy_booking_id": clean(row.get("Booking ID")), "legacy_authorization_result": clean(row.get("Authorization Result")) or None,
            "legacy_conflict_check": clean(row.get("Conflict Check")) or None, "admin_notes": clean(row.get("Admin Notes")) or None,
            "created_at": iso_datetime(row.get("Created At")),
            "metadata": {"legacy_status": original_status, "status_normalization": normalized_reason, "public_management_token_unavailable": True},
        })

    settings = []
    day_numbers = {name: index for index, name in enumerate(("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"), 1)}
    for _, row in rows_by_header(main_path, "Settings"):
        day = clean(row.get("Day"))
        if day in day_numbers:
            settings.append({"iso_day": day_numbers[day], "day_name": day, "open_time": clean(row.get("Open")), "close_time": clean(row.get("Close")), "bookable": truthy(row.get("Bookable"))})

    legacy_audit = []
    for number, row in rows_by_header(main_path, "Admin Audit"):
        if not clean(row.get("Action")):
            continue
        legacy_audit.append({"actor_display": clean(row.get("Admin")) or "Legacy system", "action": f"legacy_{clean(row.get('Action')).lower()}", "target_type": "legacy_admin_event", "result": clean(row.get("Result")).lower() or "success", "legacy_source_key": f"admin-audit:{number}", "created_at": iso_datetime(row.get("Timestamp")), "metadata": {"training": clean(row.get("Training")) or None, "notes": clean(row.get("Notes")) or None}})

    model = {
        "people": list(people.values()), "equipment": legacy_equipment, "weekly_hours": settings,
        "quiz_bank": bank, "attempts": attempts, "certifications": list(certifications_by_person_type.values()),
        "bookings": bookings, "audit": legacy_audit,
    }
    summary = {
        "mode": "include-tests" if include_tests else "exclude-tests", "people": len(model["people"]),
        "equipment": len(legacy_equipment), "weekly_hours": len(settings),
        "quiz_questions": {name: len(questions) for name, questions in bank.items()},
        "training_attempts": len(attempts), "certifications": len(model["certifications"]), "bookings": len(bookings),
        "authorized_user_rows": len(user_rows),
        "people_with_recognized_equipment_training": sum(bool(person["metadata"].get("recognized_equipment_training")) for person in people.values()),
        "certifications_by_type": dict(Counter(cert["certification_slug"] for cert in model["certifications"])),
        "certifications_by_source": dict(Counter(cert["source_kind"] for cert in model["certifications"])),
        "passed_training_attempts": sum(attempt["passed"] for attempt in attempts),
        "legacy_audit_events": len(legacy_audit), "test_rows_detected": len(test_rows),
        "test_rows_by_kind": dict(Counter(item["kind"] for item in test_rows)),
        "review_issues": len(issues), "review_issues_by_kind": dict(Counter(item["kind"] for item in issues)),
        "quiz_source_validation": quiz_validation,
        "compliance_default": "unknown_and_staff_review_required",
    }
    return model, summary, test_rows + issues


def write_reports(output_dir: Path, summary: dict[str, Any], review_rows: list[dict[str, str]]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "summary.json").write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    with (output_dir / "review.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["kind", "source", "subject"])
        writer.writeheader()
        writer.writerows(review_rows)


def apply_model(model: dict[str, Any], database_url: str) -> None:
    try:
        import psycopg
        from psycopg.types.json import Jsonb
    except ImportError as exc:
        raise RuntimeError("Install scripts/import-legacy/requirements.txt before --apply.") from exc

    with psycopg.connect(database_url) as connection, connection.transaction():
        with connection.cursor() as cursor:
            cursor.execute("set constraints all deferred")
            cursor.execute("select count(*) from public.quiz_attempts where legacy_source_key is null")
            live_attempt_count = int(cursor.fetchone()[0])
            if live_attempt_count:
                raise RuntimeError(
                    "Refusing to replace the legacy quiz bank after live quiz attempts exist. "
                    "Restore a pre-launch database or migrate quiz versions manually."
                )
            # A repeated pre-launch import rebuilds only importer-owned attempts.
            # Their answer rows cascade, allowing the private bank to be replaced
            # without weakening the foreign keys that preserve later live history.
            cursor.execute("delete from public.quiz_attempts where legacy_source_key is not null")
            for person in model["people"]:
                cursor.execute("""insert into public.people(id,email,full_name,roll_number,phone,category,organization,active,booking_privilege_active,safety_training_status,waiver_status,minor_status,migration_review_required,priority_rank,legacy_source_key,metadata)
                    values (%(id)s,%(email)s,%(full_name)s,%(roll_number)s,%(phone)s,%(category)s,%(organization)s,%(active)s,%(booking_privilege_active)s,%(safety_training_status)s,%(waiver_status)s,%(minor_status)s,%(migration_review_required)s,%(priority_rank)s,%(legacy_source_key)s,%(metadata)s)
                    on conflict (id) do update set metadata=people.metadata || excluded.metadata""", {**person, "metadata": Jsonb(person["metadata"])})

            for item in model["equipment"]:
                cursor.execute("update public.equipment set google_calendar_id=%s, status=case when %s then status else 'inactive'::public.equipment_status end, external_allowed=%s where legacy_name=%s", (item["calendar_id"], item["active"], item["external_allowed"], item["legacy_name"]))

            for hours in model["weekly_hours"]:
                cursor.execute("insert into public.weekly_hours(iso_day,day_name,open_time,close_time,bookable) values (%s,%s,%s,%s,%s) on conflict (iso_day) do update set day_name=excluded.day_name,open_time=excluded.open_time,close_time=excluded.close_time,bookable=excluded.bookable", (hours["iso_day"], hours["day_name"], hours["open_time"], hours["close_time"], hours["bookable"]))

            quiz_ids: dict[str, str] = {}
            question_ids: dict[tuple[str, str], str] = {}
            option_ids: dict[tuple[str, str, str], str] = {}
            for training, questions in model["quiz_bank"].items():
                quiz_slug = TRAINING_SLUG[training]
                cursor.execute("select id from public.quizzes where slug=%s", (quiz_slug,))
                result = cursor.fetchone()
                if not result:
                    raise RuntimeError(f"Seeded quiz {quiz_slug} is missing; apply Supabase migrations first.")
                quiz_id = str(result[0]); quiz_ids[quiz_slug] = quiz_id
                cursor.execute("update public.quizzes set active=false where id=%s", (quiz_id,))
                cursor.execute("delete from public.quiz_questions where quiz_id=%s", (quiz_id,))
                for position, question in enumerate(questions, 1):
                    question_id = stable_id("question", f"{quiz_slug}:{question['id']}")
                    question_ids[(quiz_slug, question["id"])] = question_id
                    cursor.execute("insert into public.quiz_questions(id,quiz_id,prompt,position,active,legacy_question_id) values (%s,%s,%s,%s,true,%s)", (question_id, quiz_id, question["text"], position, question["id"]))
                    for option_position, label in enumerate(question["options"], 1):
                        option_id = stable_id("option", f"{quiz_slug}:{question['id']}:{label}")
                        option_ids[(quiz_slug, question["id"], label)] = option_id
                        cursor.execute("insert into public.quiz_question_options(id,question_id,label,position,is_correct) values (%s,%s,%s,%s,%s)", (option_id, question_id, label, option_position, label == question["correct"]))
                cursor.execute("update public.quizzes set duration_minutes=8,question_count=20,pass_mark=16,active=true,version=version+1,notes='Private legacy bank imported and validated.' where id=%s", (quiz_id,))

            for attempt in model["attempts"]:
                cursor.execute("""insert into public.quiz_attempts(id,attempt_reference,quiz_id,quiz_version,participant_id,trainer_name_snapshot,status,started_at,expires_at,submitted_at,score,max_score,pass_mark,passed,legacy_source_key,source_metadata)
                    values (%(id)s,%(attempt_reference)s,%(quiz_id)s,1,%(participant_id)s,%(trainer_name_snapshot)s,%(status)s,%(started_at)s,%(expires_at)s,%(submitted_at)s,%(score)s,%(max_score)s,%(pass_mark)s,%(passed)s,%(legacy_source_key)s,%(source_metadata)s)
                    on conflict (id) do update set participant_id=excluded.participant_id,score=excluded.score,max_score=excluded.max_score,pass_mark=excluded.pass_mark,passed=excluded.passed,source_metadata=excluded.source_metadata""", {**attempt, "quiz_id": quiz_ids[attempt["quiz_slug"]], "source_metadata": Jsonb(attempt["source_metadata"])})
                cursor.execute("delete from public.quiz_attempt_answers where attempt_id=%s", (attempt["id"],))
                for answer in attempt["answers"]:
                    key = (attempt["quiz_slug"], answer["legacy_question_id"], answer["selected_label"])
                    if key not in option_ids:
                        continue
                    cursor.execute("insert into public.quiz_attempt_answers(attempt_id,question_id,selected_option_id,was_correct,answered_at) values (%s,%s,%s,%s,%s)", (attempt["id"], question_ids[(attempt["quiz_slug"], answer["legacy_question_id"])], option_ids[key], answer["was_correct"], attempt["submitted_at"]))

            cursor.execute("select slug,id from public.certification_types")
            certification_type_ids = {row[0]: str(row[1]) for row in cursor.fetchall()}
            for cert in model["certifications"]:
                cursor.execute("""insert into public.certifications(id,person_id,certification_type_id,status,issued_at,source_quiz_attempt_id,source_kind,reason,legacy_source_key,metadata)
                    values (%(id)s,%(person_id)s,%(type_id)s,%(status)s,%(issued_at)s,%(source_attempt_id)s,%(source_kind)s,%(reason)s,%(legacy_source_key)s,%(metadata)s)
                    on conflict (id) do update set source_quiz_attempt_id=excluded.source_quiz_attempt_id,metadata=certifications.metadata || excluded.metadata""", {**cert, "type_id": certification_type_ids[cert["certification_slug"]], "metadata": Jsonb(cert["metadata"])})

            cursor.execute("select legacy_name,id from public.equipment")
            equipment_ids = {row[0]: str(row[1]) for row in cursor.fetchall() if row[0]}
            for booking in model["bookings"]:
                equipment_id = equipment_ids.get(booking["legacy_equipment"])
                if not equipment_id:
                    raise RuntimeError(f"Unmapped equipment in booking: {booking['legacy_equipment']}")
                cursor.execute("""insert into public.bookings(id,booking_reference,person_id,equipment_id,status,starts_at,ends_at,contact_name,contact_email,contact_phone,contact_roll_number,contact_organization,purpose,cancelled_at,completed_at,calendar_event_id,calendar_sync_status,notification_status,legacy_booking_id,legacy_authorization_result,legacy_conflict_check,admin_notes,metadata,created_at)
                    values (%(id)s,%(booking_reference)s,%(person_id)s,%(equipment_id)s,%(status)s,%(starts_at)s,%(ends_at)s,%(contact_name)s,%(contact_email)s,%(contact_phone)s,%(contact_roll_number)s,%(contact_organization)s,%(purpose)s,%(cancelled_at)s,%(completed_at)s,%(calendar_event_id)s,%(calendar_sync_status)s,'not_configured',%(legacy_booking_id)s,%(legacy_authorization_result)s,%(legacy_conflict_check)s,%(admin_notes)s,%(metadata)s,%(created_at)s)
                    on conflict (id) do update set status=excluded.status,starts_at=excluded.starts_at,ends_at=excluded.ends_at,contact_name=excluded.contact_name,contact_email=excluded.contact_email,contact_phone=excluded.contact_phone,purpose=excluded.purpose,calendar_event_id=excluded.calendar_event_id,metadata=excluded.metadata""", {**booking, "equipment_id": equipment_id, "metadata": Jsonb(booking["metadata"])})

            for event in model["audit"]:
                cursor.execute("insert into public.audit_log(actor_display,action,target_type,result,metadata,legacy_source_key,created_at) values (%s,%s,%s,%s,%s,%s,%s) on conflict (legacy_source_key) do nothing", (event["actor_display"], event["action"], event["target_type"], event["result"], Jsonb(event["metadata"]), event["legacy_source_key"], event["created_at"]))
            cursor.execute("insert into public.audit_log(actor_display,action,target_type,metadata) values ('legacy_importer','legacy_import_completed','migration',%s)", (Jsonb({key: len(value) for key, value in model.items() if isinstance(value, list)}),))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True, help="Private extracted legacy package directory")
    parser.add_argument("--output-dir", type=Path, required=True, help="Ignored directory for non-PII reports")
    parser.add_argument("--apply", action="store_true", help="Apply within one database transaction")
    parser.add_argument("--test-records", choices=("keep", "exclude"), help="Required for --apply when tests are detected")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    include_tests = args.test_records == "keep"
    model, summary, review_rows = build_model(args.source_dir.resolve(), include_tests)
    write_reports(args.output_dir.resolve(), summary, review_rows)
    if args.apply:
        if summary["test_rows_detected"] and args.test_records is None:
            print("Refusing to apply: choose --test-records keep or --test-records exclude.", file=sys.stderr)
            return 2
        database_url = os.environ.get("DATABASE_URL", "")
        if not database_url:
            print("Refusing to apply: DATABASE_URL is not set.", file=sys.stderr)
            return 2
        apply_model(model, database_url)
        summary["applied"] = True
        write_reports(args.output_dir.resolve(), summary, review_rows)
    print(json.dumps(summary, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
